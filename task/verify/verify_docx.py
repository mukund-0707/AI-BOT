"""Check how much real content the docx extractor silently drops (tables etc.)."""

import sys
from pathlib import Path

sys.path.insert(0, r"D:\mukund\AI bot\AI bot\backend")

from docx import Document as Docx
from documents.extractors.docx import extract_docx

MEDIA = Path(r"D:\mukund\AI bot\AI bot\backend\media\documents")

for path in sorted(MEDIA.glob("*.docx")):
    print("=" * 70)
    print(path.name)

    got = extract_docx(str(path))[0]["text"]

    d = Docx(str(path))
    n_paras = len([p for p in d.paragraphs if p.text.strip()])
    n_tables = len(d.tables)

    table_chars = 0
    table_cells = 0
    sample = []
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt:
                    table_cells += 1
                    table_chars += len(txt)
                    if len(sample) < 6:
                        sample.append(txt.replace("\n", " ")[:60])

    print(f"  paragraphs kept : {n_paras}")
    print(f"  extractor chars : {len(got)}")
    print(f"  tables in file  : {n_tables}")
    print(f"  non-empty cells : {table_cells}")
    print(f"  chars INSIDE tables (currently DROPPED): {table_chars}")
    total = len(got) + table_chars
    if total:
        print(f"  >>> {table_chars / total * 100:.1f}% of document text is being LOST")
    for s in sample:
        print(f"      lost cell: {s!r}")

# PDFs: check for pages that extract to nothing (scanned/image pages)
print("=" * 70)
print("PDF empty-page check")
from documents.extractors.pdf import extract_pdf

for path in sorted(MEDIA.glob("*.pdf")):
    pages = extract_pdf(str(path))
    empty = [p["page_number"] for p in pages if not p["text"].strip()]
    chars = sum(len(p["text"]) for p in pages)
    print(
        f"  {path.name}: {len(pages)} pages, {chars} chars, empty pages={empty or 'none'}"
    )
